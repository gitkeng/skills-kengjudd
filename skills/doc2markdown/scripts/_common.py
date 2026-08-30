"""Shared helpers for the doc2markdown extraction scripts.

Imported by probe.py / extract_pdf.py / extract_docx.py / extract_xlsx.py /
merge.py.  When a script is launched with `uv run <script>`, the script's own
directory is sys.path[0], so a plain `import _common` resolves to this file.

Everything here is deterministic string/grid plumbing.  No model, no guessing.
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

# --------------------------------------------------------------------------
# names
# --------------------------------------------------------------------------

_UNSAFE = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def base_name(path) -> str:
    """Stem of a path, NFC-normalised (matters for Thai filenames on macOS)."""
    return unicodedata.normalize("NFC", Path(path).stem)


def slug(name: str, fallback: str = "item") -> str:
    """Filesystem-safe token.  Keeps Thai/CJK characters, kills separators."""
    s = unicodedata.normalize("NFC", str(name)).strip()
    s = _UNSAFE.sub("_", s)
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"_{2,}", "_", s).strip("._")
    return s or fallback


# --------------------------------------------------------------------------
# markdown tables
# --------------------------------------------------------------------------


def cell(value) -> str:
    """Render one value as a markdown table cell.

    Escapes backslashes first, then pipes, then folds newlines to <br> so a
    multi-line cell cannot break the row.
    """
    if value is None:
        return ""
    s = str(value)
    if not s.strip():
        return ""
    s = s.replace("\\", "\\\\").replace("|", "\\|")
    s = re.sub(r"\r\n|\r|\n", "<br>", s)
    s = re.sub(r"[ \t]{2,}", " ", s)
    return s.strip()


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def to_md_table(rows, header_rows: int = 1) -> str:
    """Grid -> GitHub-flavoured markdown table.

    header_rows=0 emits an empty header row (markdown requires one).
    header_rows>1 folds the header band into a single row joined with " / ",
    dropping duplicates so a merged header does not repeat itself.
    """
    rows = [list(r) for r in rows if r is not None]
    if not rows:
        return ""
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return ""
    grid = [r + [None] * (width - len(r)) for r in rows]

    header_rows = max(0, min(header_rows, len(grid)))
    out = []

    if header_rows == 0:
        out.append("|" + "|".join(" " * 3 for _ in range(width)) + "|")
        body = grid
    else:
        band = grid[:header_rows]
        head = []
        for col in range(width):
            parts = []
            for row in band:
                v = cell(row[col])
                if v and v not in parts:
                    parts.append(v)
            head.append(" / ".join(parts))
        out.append("| " + " | ".join(head) + " |")
        body = grid[header_rows:]

    out.append("|" + "|".join(" --- " for _ in range(width)) + "|")
    for row in body:
        out.append("| " + " | ".join(cell(v) for v in row) + " |")
    return "\n".join(out)


def trim_grid(rows):
    """Drop all-blank leading/trailing rows and leading/trailing columns.

    Interior blank rows/columns are preserved -- in a spreadsheet they often
    separate two logically distinct tables, and dropping them would silently
    merge unrelated data.
    """
    grid = [list(r) for r in rows]
    while grid and all(is_blank(v) for v in grid[0]):
        grid.pop(0)
    while grid and all(is_blank(v) for v in grid[-1]):
        grid.pop()
    if not grid:
        return []

    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]

    first = 0
    while first < width and all(is_blank(r[first]) for r in grid):
        first += 1
    last = width - 1
    while last >= first and all(is_blank(r[last]) for r in grid):
        last -= 1
    if first > last:
        return []
    return [r[first : last + 1] for r in grid]


def blank_column_indexes(rows) -> list[int]:
    """Indexes of fully-blank interior columns, for reporting in the sidecar."""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    out = []
    for i in range(width):
        if all(is_blank(r[i]) if i < len(r) else True for r in rows):
            out.append(i)
    return out


# --------------------------------------------------------------------------
# text quality
# --------------------------------------------------------------------------

_PUA = re.compile("[\ue000-\uf8ff\ufffd]")


def garbled_ratio(text: str) -> float:
    """Fraction of characters that are replacement / private-use glyphs.

    A PDF with a broken ToUnicode CMap decodes into private-use codepoints.
    A high ratio means the text layer is unusable and the page must be read
    from its rendered image instead.
    """
    if not text:
        return 0.0
    return len(_PUA.findall(text)) / len(text)


# --------------------------------------------------------------------------
# pptx shape walking (shared by probe.py and extract_pptx.py)
# --------------------------------------------------------------------------

PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def group_transform(group):
    """((off), (ext), (chOff), (chExt)) mapping a group's child space to the
    parent space, or None when the xfrm is missing or degenerate.

    Shapes inside a group report positions in the group's *child* coordinate
    system -- using them raw puts every grouped shape in the wrong place, and
    visual reading order silently breaks.
    """
    grp_sp_pr = group._element.find(PML + "grpSpPr")
    if grp_sp_pr is None:
        return None
    xfrm = grp_sp_pr.find(DML + "xfrm")
    if xfrm is None:
        return None

    def point(tag):
        el = xfrm.find(DML + tag)
        return (int(el.get("x")), int(el.get("y"))) if el is not None else None

    def extent(tag):
        el = xfrm.find(DML + tag)
        return (int(el.get("cx")), int(el.get("cy"))) if el is not None else None

    off, ext, ch_off, ch_ext = point("off"), extent("ext"), point("chOff"), extent("chExt")
    if None in (off, ext, ch_off, ch_ext) or not ch_ext[0] or not ch_ext[1]:
        return None
    return off, ext, ch_off, ch_ext


def iter_shapes(shapes, ox=0.0, oy=0.0, sx=1.0, sy=1.0):
    """Yield (shape, abs_left, abs_top, abs_width, abs_height) in slide space,
    recursing into groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        left = ox + (shape.left or 0) * sx
        top = oy + (shape.top or 0) * sy
        width = (shape.width or 0) * sx
        height = (shape.height or 0) * sy
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            transform = group_transform(shape)
            if transform is None:
                yield from iter_shapes(shape.shapes, left, top, sx, sy)
            else:
                (off, ext, ch_off, ch_ext) = transform
                child_sx = sx * ext[0] / ch_ext[0]
                child_sy = sy * ext[1] / ch_ext[1]
                child_ox = ox + (off[0] - ch_off[0] * ext[0] / ch_ext[0]) * sx
                child_oy = oy + (off[1] - ch_off[1] * ext[1] / ch_ext[1]) * sy
                yield from iter_shapes(shape.shapes, child_ox, child_oy, child_sx, child_sy)
        else:
            yield shape, left, top, width, height


# --------------------------------------------------------------------------
# io
# --------------------------------------------------------------------------


def write_json(path, data) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_text(path, text: str) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")


def emit(data) -> None:
    """Print a JSON result on stdout (the script's machine-readable return)."""
    print(json.dumps(data, ensure_ascii=False, indent=2))
