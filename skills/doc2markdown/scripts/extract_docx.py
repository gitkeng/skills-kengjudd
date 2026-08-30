#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1"]
# ///
"""Turn a .docx into ordered markdown chunks plus its images.

Walks the document body in true XML order so paragraphs and tables interleave
exactly as they do in Word -- iterating `document.paragraphs` then
`document.tables` (the obvious approach) silently reorders the whole document.

Emits per chunk a markdown draft and the matching plain text, so a reviewing
agent can check the draft against the source without re-reading the .docx.

    uv run extract_docx.py <file.docx> --work <workdir> --assets <assetsdir> [--base NAME]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from urllib.parse import quote

import _common as C

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
V = "{urn:schemas-microsoft-com:vml}"

DEFAULT_CHUNK_CHARS = 6000


# --------------------------------------------------------------------------
# images
# --------------------------------------------------------------------------


class ImageBank:
    """Saves each embedded image once, numbered by first appearance."""

    def __init__(self, part, assets_dir: Path, base: str, min_bytes: int):
        self.part = part
        self.assets_dir = assets_dir
        self.base = base
        self.min_bytes = min_bytes
        self.by_rid: dict[str, dict] = {}
        self.by_digest: dict[str, dict] = {}
        self.saved: list[dict] = []

    def resolve(self, rid: str) -> dict | None:
        if rid in self.by_rid:
            return self.by_rid[rid]
        try:
            image_part = self.part.related_parts[rid]
            blob = image_part.blob
        except Exception:
            return None
        if not blob or len(blob) < self.min_bytes:
            return None

        import hashlib

        digest = hashlib.sha1(blob).hexdigest()
        entry = self.by_digest.get(digest)
        if entry is None:
            ext = Path(str(image_part.partname)).suffix.lower().lstrip(".") or "png"
            index = len(self.saved) + 1
            name = f"{self.base}_{index:02d}.{ext}"
            self.assets_dir.mkdir(parents=True, exist_ok=True)
            (self.assets_dir / name).write_bytes(blob)
            rel = f"{self.assets_dir.name}/{name}"
            entry = {
                "index": index,
                "file": name,
                "rel_path": rel,
                "md": f"![]({quote(rel)})",
                "bytes": len(blob),
            }
            self.by_digest[digest] = entry
            self.saved.append(entry)
        self.by_rid[rid] = entry
        return entry


# --------------------------------------------------------------------------
# inline content
# --------------------------------------------------------------------------


def _run_text(run_el) -> str:
    out = []
    for node in run_el.iter():
        tag = node.tag
        if tag == W + "t":
            out.append(node.text or "")
        elif tag == W + "tab":
            out.append("\t")
        elif tag in (W + "br", W + "cr"):
            out.append("\n")
        elif tag == W + "noBreakHyphen":
            out.append("-")
    return "".join(out)


def _is_on(rpr, name: str) -> bool:
    if rpr is None:
        return False
    el = rpr.find(W + name)
    if el is None:
        return False
    val = el.get(W + "val")
    return val not in ("0", "false", "off")


def _run_images(run_el, bank: ImageBank) -> list[dict]:
    found = []
    for blip in run_el.iter(A + "blip"):
        rid = blip.get(R + "embed") or blip.get(R + "link")
        if rid:
            entry = bank.resolve(rid)
            if entry:
                found.append(entry)
    for data in run_el.iter(V + "imagedata"):
        rid = data.get(R + "id")
        if rid:
            entry = bank.resolve(rid)
            if entry:
                found.append(entry)
    return found


def _textbox_paragraphs(run_el) -> list[str]:
    """Text trapped inside a Word text box, which normal walks never reach."""
    out = []
    for box in run_el.iter(W + "txbxContent"):
        for para in box.iter(W + "p"):
            text = "".join(
                node.text or "" for node in para.iter(W + "t")
            ).strip()
            if text:
                out.append(text)
    return out


def render_inline(p_el, bank: ImageBank, part) -> tuple[str, str, list[dict], list[str]]:
    """Return (markdown, plain, images, textboxes) for one paragraph."""
    md_parts: list[str] = []
    plain_parts: list[str] = []
    images: list[dict] = []
    boxes: list[str] = []

    def emit_run(run_el, link: str | None = None):
        text = _run_text(run_el)
        images.extend(_run_images(run_el, bank))
        boxes.extend(_textbox_paragraphs(run_el))
        if not text:
            return
        plain_parts.append(text)
        rpr = run_el.find(W + "rPr")
        body = text
        stripped = body.strip()
        if stripped:
            lead = body[: len(body) - len(body.lstrip())]
            trail = body[len(body.rstrip()) :]
            if _is_on(rpr, "b"):
                stripped = f"**{stripped}**"
            if _is_on(rpr, "i"):
                stripped = f"*{stripped}*"
            body = lead + stripped + trail
        if link:
            body = f"[{body.strip()}]({link})"
        md_parts.append(body)

    for child in p_el.iterchildren():
        tag = child.tag
        if tag == W + "r":
            emit_run(child)
        elif tag == W + "hyperlink":
            url = None
            rid = child.get(R + "id")
            if rid:
                try:
                    url = part.rels[rid].target_ref
                except Exception:
                    url = None
            for run_el in child.findall(W + "r"):
                emit_run(run_el, url)
        elif tag == W + "ins":  # tracked insertion -- accepted text
            for run_el in child.findall(W + "r"):
                emit_run(run_el)
        # w:del (tracked deletion) is intentionally skipped: it is not in the
        # document as it stands.

    return "".join(md_parts), "".join(plain_parts), images, boxes


# --------------------------------------------------------------------------
# numbering
# --------------------------------------------------------------------------


def build_numbering(document) -> dict:
    """numId -> {level: 'ordered'|'bullet'} from numbering.xml."""
    mapping: dict = {}
    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return mapping

    abstract: dict = {}
    for anum in numbering.findall(W + "abstractNum"):
        aid = anum.get(W + "abstractNumId")
        levels = {}
        for lvl in anum.findall(W + "lvl"):
            ilvl = lvl.get(W + "ilvl")
            fmt_el = lvl.find(W + "numFmt")
            fmt = fmt_el.get(W + "val") if fmt_el is not None else "bullet"
            levels[ilvl] = "bullet" if fmt in ("bullet", "none") else "ordered"
        abstract[aid] = levels

    for num in numbering.findall(W + "num"):
        nid = num.get(W + "numId")
        ref = num.find(W + "abstractNumId")
        if ref is not None:
            mapping[nid] = abstract.get(ref.get(W + "val"), {})
    return mapping


def list_info(p_el, numbering: dict):
    ppr = p_el.find(W + "pPr")
    if ppr is None:
        return None
    numpr = ppr.find(W + "numPr")
    if numpr is None:
        return None
    ilvl_el = numpr.find(W + "ilvl")
    numid_el = numpr.find(W + "numId")
    ilvl = ilvl_el.get(W + "val") if ilvl_el is not None else "0"
    numid = numid_el.get(W + "val") if numid_el is not None else None
    if numid in (None, "0"):
        return None
    kind = numbering.get(numid, {}).get(ilvl, "bullet")
    try:
        depth = int(ilvl)
    except (TypeError, ValueError):
        depth = 0
    return {"kind": kind, "depth": depth}


def heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if name in ("title", "หัวเรื่อง"):
        return 1
    for prefix in ("heading ", "heading"):
        if name.startswith(prefix):
            tail = name[len(prefix) :].strip()
            if tail.isdigit():
                return min(int(tail) + 1, 6)
    return None


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------


def cell_text(cell) -> str:
    lines = []
    for para in cell.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for nested in cell.tables:
        for row in nested.rows:
            lines.append(" / ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def table_grid(table) -> list[list[str]]:
    """Grid with merged-cell continuations blanked instead of duplicated."""
    grid = []
    seen: set[int] = set()
    for row in table.rows:
        out_row = []
        try:
            cells = row.cells
        except (IndexError, ValueError):
            continue
        for c in cells:
            key = id(c._tc)
            if key in seen:
                out_row.append("")  # continuation of a merged cell
            else:
                seen.add(key)
                out_row.append(cell_text(c))
        grid.append(out_row)
    return grid


# --------------------------------------------------------------------------
# body walk
# --------------------------------------------------------------------------


def walk_body(document, bank: ImageBank) -> list[dict]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    numbering = build_numbering(document)
    part = document.part
    blocks: list[dict] = []

    for child in document.element.body.iterchildren():
        tag = child.tag
        if tag == W + "p":
            para = Paragraph(child, document)
            md, plain, images, boxes = render_inline(child, bank, part)
            level = heading_level(para.style.name if para.style is not None else "")
            listing = list_info(child, numbering)

            for entry in images:
                blocks.append({"type": "image", "md": entry["md"], "plain": "", "image": entry})
            for box in boxes:
                blocks.append({"type": "textbox", "md": f"> {box}", "plain": box})

            text = md.strip()
            if not text:
                continue
            if level:
                blocks.append({"type": "heading", "level": level, "md": f"{'#' * level} {text}", "plain": plain})
            elif listing:
                indent = "  " * listing["depth"]
                marker = "1." if listing["kind"] == "ordered" else "-"
                blocks.append({"type": "list", "md": f"{indent}{marker} {text}", "plain": plain})
            else:
                blocks.append({"type": "paragraph", "md": text, "plain": plain})

        elif tag == W + "tbl":
            table = Table(child, document)
            grid = C.trim_grid(table_grid(table))
            if not grid:
                continue
            md = C.to_md_table(grid, header_rows=1)
            plain = "\n".join(" ".join(str(v or "") for v in row) for row in grid)
            blocks.append(
                {
                    "type": "table",
                    "md": md,
                    "plain": plain,
                    "rows": len(grid),
                    "cols": max(len(r) for r in grid),
                }
            )

            for cell_el in child.iter(W + "tc"):
                for run_el in cell_el.iter(W + "r"):
                    for entry in _run_images(run_el, bank):
                        blocks.append(
                            {"type": "image", "md": entry["md"], "plain": "", "image": entry}
                        )

    return blocks


def section_furniture(document) -> dict:
    """Header/footer text -- page furniture, kept out of the body by default."""
    out = {"headers": [], "footers": []}
    for section in document.sections:
        for key, part in (("headers", section.header), ("footers", section.footer)):
            try:
                text = "\n".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:
                continue
            if text and text not in out[key]:
                out[key].append(text)
    return out


def chunk_blocks(blocks: list[dict], limit: int) -> list[list[dict]]:
    """Split at headings where possible; never inside a table."""
    chunks: list[list[dict]] = []
    current: list[dict] = []
    size = 0

    for block in blocks:
        length = len(block.get("md", ""))
        starts_section = block["type"] == "heading" and block.get("level", 9) <= 2
        if current and (
            (size >= limit and starts_section)
            or (size >= limit * 2 and block["type"] != "table")
        ):
            chunks.append(current)
            current, size = [], 0
        current.append(block)
        size += length

    if current:
        chunks.append(current)
    return chunks or [[]]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--work", required=True)
    ap.add_argument("--assets", required=True)
    ap.add_argument("--base", default=None)
    ap.add_argument("--chunk-chars", type=int, default=DEFAULT_CHUNK_CHARS)
    ap.add_argument("--min-image-bytes", type=int, default=3000)
    args = ap.parse_args()

    import docx

    path = Path(args.file).expanduser()
    if not path.exists():
        C.emit({"ok": False, "error": f"file not found: {path}"})
        return 1
    if path.suffix.lower() != ".docx":
        C.emit({"ok": False, "error": "only .docx is supported; convert .doc first"})
        return 2

    base = args.base or C.base_name(path)
    work = Path(args.work).expanduser()
    assets_dir = Path(args.assets).expanduser()

    document = docx.Document(str(path))
    bank = ImageBank(document.part, assets_dir, base, args.min_image_bytes)
    blocks = walk_body(document, bank)
    furniture = section_furniture(document)
    chunks = chunk_blocks(blocks, args.chunk_chars)

    chunk_meta = []
    for index, chunk in enumerate(chunks, start=1):
        md = "\n\n".join(b["md"] for b in chunk if b.get("md"))
        plain = "\n".join(b["plain"] for b in chunk if b.get("plain"))
        md_rel = f"chunks/c{index:02d}.md"
        txt_rel = f"chunks/c{index:02d}.txt"
        C.write_text(work / md_rel, md + "\n")
        C.write_text(work / txt_rel, plain + "\n")
        chunk_meta.append(
            {
                "index": index,
                "draft_file": md_rel,
                "text_file": txt_rel,
                "chars": len(md),
                "blocks": len(chunk),
                "tables": sum(1 for b in chunk if b["type"] == "table"),
                "images": [b["image"]["file"] for b in chunk if b["type"] == "image"],
                "headings": [b["md"] for b in chunk if b["type"] == "heading"],
            }
        )

    manifest = {
        "ok": True,
        "kind": "docx",
        "source": str(path.resolve()),
        "base": base,
        "work_dir": str(work.resolve()),
        "assets_dir": str(assets_dir.resolve()),
        "assets_dir_name": assets_dir.name,
        "block_count": len(blocks),
        "assets": bank.saved,
        "headers_footers": furniture,
        "chunks": chunk_meta,
        "summary": {
            "chunk_count": len(chunk_meta),
            "asset_count": len(bank.saved),
            "table_count": sum(1 for b in blocks if b["type"] == "table"),
            "heading_count": sum(1 for b in blocks if b["type"] == "heading"),
            "textbox_count": sum(1 for b in blocks if b["type"] == "textbox"),
        },
    }
    C.write_json(work / "manifest.json", manifest)
    C.emit(manifest)
    return 0


if __name__ == "__main__":
    sys.exit(main())
