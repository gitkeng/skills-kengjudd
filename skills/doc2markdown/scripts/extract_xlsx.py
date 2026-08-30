#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Turn a workbook into one markdown table per sheet, plus its images.

A spreadsheet is a grid, not prose, so nothing here has to be guessed at the
paragraph level -- the whole job is getting the *values* right:

  * a merged cell is written into every position it covers, so each markdown
    row stands on its own;  markdown has no rowspan, and a blanked
    continuation row loses which group it belonged to.
  * cached formula results are preferred; when the workbook was never opened by
    Excel the cache is empty and the formula itself is written instead, marked
    so nobody reads `=SUM(B2:B9)` as a number.
  * dates, percentages and booleans are rendered the way the sheet displays
    them, not as the serial numbers openpyxl hands back.

Hidden rows/columns are extracted like any other -- they are still data -- but
reported in the manifest so a reviewer knows they were not visible on screen.

    uv run extract_xlsx.py <file.xlsx> --work <workdir> --assets <assetsdir> [--base NAME]

Writes <workdir>/manifest.json and prints it on stdout.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import sys
from pathlib import Path
from urllib.parse import quote

import _common as C

# Sheets wider than this stop being readable as markdown; the raw TSV stays the
# authority and the manifest says so.
WIDE_SHEET_COLS = 25

# Leading rows to consider when guessing how deep the header band goes.
MAX_HEADER_ROWS = 3


class Cell:
    """One rendered cell.

    Carries its display text plus the facts the header guess and the manifest
    need.  `__str__` is the display text, so _common's grid helpers (which call
    str()) work on these unchanged.
    """

    __slots__ = ("text", "numeric", "formula")

    def __init__(self, text: str = "", numeric: bool = False, formula: str | None = None):
        self.text = text
        self.numeric = numeric
        self.formula = formula

    def __str__(self) -> str:
        return self.text


BLANK = Cell()


# --------------------------------------------------------------------------
# value rendering
# --------------------------------------------------------------------------


def _fmt_number(value) -> str:
    """Plain decimal, no thousands separators, no float noise."""
    if isinstance(value, int):
        return str(value)
    text = f"{value:.10g}"
    if "e" in text or "E" in text:  # very large/small: keep full precision
        return repr(value)
    return text


def render_value(value, number_format: str) -> tuple[str, bool]:
    """(display text, is_numeric) for one cell value."""
    if value is None:
        return "", False
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), False
    if isinstance(value, dt.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat(), False
        return value.isoformat(sep=" "), False
    if isinstance(value, dt.date):
        return value.isoformat(), False
    if isinstance(value, dt.time):
        return value.isoformat(), False
    if isinstance(value, dt.timedelta):
        return str(value), False
    if isinstance(value, (int, float)):
        if "%" in (number_format or ""):
            return _fmt_number(value * 100) + "%", True
        return _fmt_number(value), True
    return str(value).strip(), False


def build_cell(data_cell, formula_cell) -> Cell:
    """Render one position, falling back to its formula when the cache is cold."""
    value = data_cell.value if data_cell is not None else None
    formula = None
    if formula_cell is not None and isinstance(formula_cell.value, str):
        if formula_cell.value.startswith("="):
            formula = formula_cell.value

    number_format = getattr(data_cell, "number_format", "") or ""
    text, numeric = render_value(value, number_format)

    if not text and formula:
        # Never opened in Excel, so no cached result exists.  Showing the
        # formula is honest; showing nothing would read as an empty cell.
        return Cell(formula, False, formula)

    link = getattr(data_cell, "hyperlink", None)
    target = getattr(link, "target", None) if link is not None else None
    if target and text:
        text = f"[{text}]({quote(str(target), safe=':/?#[]@!$&()*+,;=%~')})"

    return Cell(text, numeric, formula)


# --------------------------------------------------------------------------
# grid
# --------------------------------------------------------------------------


def sheet_grid(ws_data, ws_formula) -> list[list[Cell]]:
    """Full used range as Cells, with merged values spread over their range."""
    rows = ws_data.max_row or 0
    cols = ws_data.max_column or 0
    if rows == 0 or cols == 0:
        return []

    grid = [[BLANK] * cols for _ in range(rows)]
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            data_cell = ws_data.cell(row=r, column=c)
            formula_cell = ws_formula.cell(row=r, column=c) if ws_formula is not None else None
            grid[r - 1][c - 1] = build_cell(data_cell, formula_cell)

    for rng in ws_data.merged_cells.ranges:
        anchor = grid[rng.min_row - 1][rng.min_col - 1]
        if not anchor.text:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[r - 1][c - 1] = anchor

    return grid


def trim_offsets(grid) -> tuple[int, int, int, int] | None:
    """(first_row, last_row, first_col, last_col) of the non-blank envelope."""
    if not grid:
        return None
    height = len(grid)
    width = max(len(r) for r in grid)

    top = 0
    while top < height and all(C.is_blank(v) for v in grid[top]):
        top += 1
    bottom = height - 1
    while bottom >= top and all(C.is_blank(v) for v in grid[bottom]):
        bottom -= 1
    if top > bottom:
        return None

    left = 0
    while left < width and all(
        C.is_blank(grid[r][left]) if left < len(grid[r]) else True for r in range(top, bottom + 1)
    ):
        left += 1
    right = width - 1
    while right >= left and all(
        C.is_blank(grid[r][right]) if right < len(grid[r]) else True for r in range(top, bottom + 1)
    ):
        right -= 1
    if left > right:
        return None
    return top, bottom, left, right


def collapse_banner_rows(grid) -> list[int]:
    """Un-spread rows that are a single merge across the whole width.

    Spreading a merge is right for a group column, where every row has to stand
    on its own.  A full-width merge is not that -- it is a section title banner,
    and repeating it once per column is noise.  Markdown has no colspan, so the
    text stays in the first cell and the rest go blank.

    Merged positions share one Cell object, so identity settles this exactly.
    Returns the indexes of the rows it touched.
    """
    banners = []
    for index, row in enumerate(grid):
        if len(row) < 2 or C.is_blank(row[0]):
            continue
        if all(c is row[0] for c in row):
            grid[index] = [row[0]] + [BLANK] * (len(row) - 1)
            banners.append(index)
    return banners


def guess_header_rows(grid, banners: list[int]) -> int:
    """How many leading rows form the header band.

    A header row is text-only; the data underneath is what carries numbers.  So
    the band is the run of leading rows with no numeric cell -- but only if some
    row below actually has one, otherwise the sheet is all text and the first
    row is the header by convention.  A banner row ends the band: it belongs to
    the body, labelling the rows beneath it.
    """
    if not grid:
        return 0
    has_numbers_below = any(any(c.numeric for c in row) for row in grid[1:])
    if not has_numbers_below:
        return 1

    banner_set = set(banners)
    count = 0
    for index, row in enumerate(grid[:MAX_HEADER_ROWS]):
        if index in banner_set:
            break
        if any(c.numeric for c in row):
            break
        if all(C.is_blank(c) for c in row):
            break
        count += 1
    return max(1, count)


# --------------------------------------------------------------------------
# images
# --------------------------------------------------------------------------


def collect_images(ws, assets_dir: Path, base: str, sheet_index: int, seen: dict, saved: list):
    """Save each embedded picture once, numbered across the whole workbook."""
    out = []
    for image in getattr(ws, "_images", []) or []:
        try:
            blob = image._data()
        except Exception as exc:
            print(f"[warn] sheet {sheet_index}: image unreadable: {exc}", file=sys.stderr)
            continue
        if not blob:
            continue

        digest = hashlib.sha1(blob).hexdigest()
        entry = seen.get(digest)
        if entry is None:
            ext = (getattr(image, "format", None) or "png").lower().lstrip(".")
            index = len(saved) + 1
            name = f"{base}_{index:02d}.{ext}"
            assets_dir.mkdir(parents=True, exist_ok=True)
            (assets_dir / name).write_bytes(blob)
            rel = f"{assets_dir.name}/{name}"
            entry = {
                "index": index,
                "file": name,
                "rel_path": rel,
                "md": f"![]({quote(rel)})",
                "bytes": len(blob),
                "sheets": [],
            }
            seen[digest] = entry
            saved.append(entry)
        if sheet_index not in entry["sheets"]:
            entry["sheets"].append(sheet_index)
        out.append(entry["file"])
    return out


def anchor_ref(obj) -> str | None:
    """Best-effort 'where on the sheet' for a chart or image."""
    anchor = getattr(obj, "anchor", None)
    if isinstance(anchor, str):
        return anchor
    frm = getattr(anchor, "_from", None)
    if frm is None:
        return None
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(frm.col + 1)}{frm.row + 1}"


# --------------------------------------------------------------------------
# output
# --------------------------------------------------------------------------


def write_parts(work: Path, grid, header_rows: int, max_rows: int, sheet_index: int) -> list[dict]:
    """Markdown for one sheet, split into parts that each repeat the header."""
    band = grid[:header_rows]
    body = grid[header_rows:]
    if not body:
        body = [[BLANK] * (len(grid[0]) if grid else 1)]

    parts = []
    total = len(body)
    for start in range(0, total, max_rows):
        block = body[start : start + max_rows]
        index = len(parts) + 1
        rel = f"sheets/s{sheet_index:02d}_p{index:02d}.md"
        md = C.to_md_table(band + block, header_rows=header_rows)
        if total > max_rows:
            md = (
                f"<!-- sheet {sheet_index} rows {start + header_rows + 1}"
                f"-{start + header_rows + len(block)} of {total + header_rows} -->\n{md}"
            )
        C.write_text(work / rel, md + "\n")
        parts.append(
            {
                "index": index,
                "draft_file": rel,
                "first_body_row": start + header_rows + 1,
                "last_body_row": start + header_rows + len(block),
                "rows": len(block),
            }
        )
    return parts


def write_tsv(work: Path, grid, sheet_index: int) -> str:
    """Raw grid, tab-separated -- the authority when a table is too wide."""
    rel = f"sheets/s{sheet_index:02d}.tsv"
    lines = []
    for row in grid:
        lines.append("\t".join(str(c).replace("\t", " ").replace("\n", " ") for c in row))
    C.write_text(work / rel, "\n".join(lines) + "\n")
    return rel


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--work", required=True)
    ap.add_argument("--assets", required=True)
    ap.add_argument("--base", default=None)
    ap.add_argument("--max-rows", type=int, default=200, help="body rows per markdown part")
    ap.add_argument("--include-hidden-sheets", action="store_true")
    args = ap.parse_args()

    import openpyxl

    path = Path(args.file).expanduser()
    if not path.exists():
        C.emit({"ok": False, "error": f"file not found: {path}"})
        return 1
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        C.emit({"ok": False, "error": "only .xlsx/.xlsm is supported; convert .xls first"})
        return 2

    base = args.base or C.base_name(path)
    work = Path(args.work).expanduser()
    assets_dir = Path(args.assets).expanduser()

    try:
        wb_data = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        C.emit({"ok": False, "error": f"{type(exc).__name__}: {exc}", "path": str(path)})
        return 3
    try:
        wb_formula = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        wb_formula = None  # formulas are a bonus, not a requirement

    seen_images: dict = {}
    saved_images: list = []
    sheets = []
    skipped = []

    for index, ws in enumerate(wb_data.worksheets, start=1):
        if ws.sheet_state != "visible" and not args.include_hidden_sheets:
            skipped.append({"index": index, "name": ws.title, "state": ws.sheet_state})
            continue

        ws_formula = None
        if wb_formula is not None and ws.title in wb_formula.sheetnames:
            ws_formula = wb_formula[ws.title]

        full = sheet_grid(ws, ws_formula)
        box = trim_offsets(full)
        images = collect_images(ws, assets_dir, base, index, seen_images, saved_images)
        charts = [
            {"title": str(getattr(ch, "title", "") or ""), "at": anchor_ref(ch)}
            for ch in (getattr(ws, "_charts", []) or [])
        ]

        if box is None:
            sheets.append(
                {
                    "index": index,
                    "name": ws.title,
                    "slug": C.slug(ws.title, f"sheet{index}"),
                    "empty": True,
                    "rows": 0,
                    "cols": 0,
                    "parts": [],
                    "images": images,
                    "charts": charts,
                }
            )
            continue

        top, bottom, left, right = box
        grid = [row[left : right + 1] for row in full[top : bottom + 1]]
        banners = collapse_banner_rows(grid)
        header_rows = guess_header_rows(grid, banners)

        from openpyxl.utils import get_column_letter

        hidden_rows = [
            r
            for r in range(top + 1, bottom + 2)
            if (d := ws.row_dimensions.get(r)) is not None and d.hidden
        ]
        hidden_cols = [
            get_column_letter(c)
            for c in range(left + 1, right + 2)
            if (d := ws.column_dimensions.get(get_column_letter(c))) is not None and d.hidden
        ]
        comments = sum(
            1
            for row in ws.iter_rows(min_row=top + 1, max_row=bottom + 1)
            for cl in row
            if getattr(cl, "comment", None) is not None
        )

        parts = write_parts(work, grid, header_rows, args.max_rows, index)
        tsv = write_tsv(work, grid, index)

        sheets.append(
            {
                "index": index,
                "name": ws.title,
                "slug": C.slug(ws.title, f"sheet{index}"),
                "empty": False,
                "rows": len(grid),
                "cols": len(grid[0]),
                "origin": f"{get_column_letter(left + 1)}{top + 1}",
                "header_rows": header_rows,
                "banner_rows": [i + 1 for i in banners],
                "tsv_file": tsv,
                "parts": parts,
                "merged_ranges": len(ws.merged_cells.ranges),
                "blank_interior_columns": C.blank_column_indexes(grid),
                "formula_cells": sum(1 for row in grid for c in row if c.formula),
                "uncached_formulas": sum(
                    1 for row in grid for c in row if c.formula and c.text == c.formula
                ),
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_cols,
                "comments": comments,
                "images": images,
                "charts": charts,
                "too_wide_for_markdown": len(grid[0]) > WIDE_SHEET_COLS,
            }
        )

    wb_data.close()
    if wb_formula is not None:
        wb_formula.close()

    manifest = {
        "ok": True,
        "kind": "xlsx",
        "source": str(path.resolve()),
        "base": base,
        "work_dir": str(work.resolve()),
        "assets_dir": str(assets_dir.resolve()),
        "assets_dir_name": assets_dir.name,
        "sheet_count": len(sheets),
        "assets": saved_images,
        "skipped_sheets": skipped,
        "sheets": sheets,
        "summary": {
            "part_count": sum(len(s["parts"]) for s in sheets),
            "asset_count": len(saved_images),
            "empty_sheets": [s["name"] for s in sheets if s["empty"]],
            "wide_sheets": [s["name"] for s in sheets if s.get("too_wide_for_markdown")],
            "sheets_with_uncached_formulas": [
                s["name"] for s in sheets if s.get("uncached_formulas")
            ],
            "sheets_with_hidden_data": [
                s["name"] for s in sheets if s.get("hidden_rows") or s.get("hidden_columns")
            ],
            "sheets_with_charts": [s["name"] for s in sheets if s.get("charts")],
        },
    }
    C.write_json(work / "manifest.json", manifest)
    C.emit(manifest)
    return 0


if __name__ == "__main__":
    sys.exit(main())
